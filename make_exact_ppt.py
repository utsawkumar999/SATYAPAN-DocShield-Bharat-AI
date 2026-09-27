# -*- coding: utf-8 -*-
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

BG_DARK = RGBColor(10, 15, 29)
CARD_BG = RGBColor(18, 26, 47)
CARD_INNER = RGBColor(13, 20, 38)
ACCENT_BLUE = RGBColor(59, 130, 246)
ACCENT_CYAN = RGBColor(6, 182, 212)
ACCENT_GREEN = RGBColor(16, 185, 129)
ACCENT_AMBER = RGBColor(245, 158, 11)
ACCENT_RED = RGBColor(239, 68, 68)
TEXT_WHITE = RGBColor(248, 250, 252)
TEXT_MUTED = RGBColor(148, 163, 184)
BORDER_COLOR = RGBColor(30, 41, 59)

def set_bg(slide):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_DARK
    bg.line.fill.background()
    return bg

def add_header(slide, title_text, slide_num=1):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(0.35), Inches(12.133), Inches(0.85))
    bar.fill.solid()
    bar.fill.fore_color.rgb = CARD_BG
    bar.line.color.rgb = BORDER_COLOR
    tf = bar.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.25)
    tf.margin_top = Inches(0.1)
    
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE
    p.font.name = 'Arial'

    tag_box = slide.shapes.add_textbox(Inches(8.2), Inches(0.42), Inches(4.3), Inches(0.7))
    tf_tag = tag_box.text_frame
    p_tag = tf_tag.paragraphs[0]
    p_tag.alignment = PP_ALIGN.RIGHT
    p_tag.text = f'Team: Digital Warriors  |  Slide {slide_num}/6'
    p_tag.font.size = Pt(11)
    p_tag.font.bold = True
    p_tag.font.color.rgb = ACCENT_CYAN
    p_tag.font.name = 'Arial'

# SLIDE 1
slide1 = prs.slides.add_slide(blank_layout)
set_bg(slide1)
card1 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(0.6), Inches(12.133), Inches(6.3))
card1.fill.solid()
card1.fill.fore_color.rgb = CARD_BG
card1.line.color.rgb = ACCENT_BLUE
card1.line.width = Pt(2)
tf1 = card1.text_frame
tf1.word_wrap = True
tf1.margin_left = Inches(0.5)
tf1.margin_top = Inches(0.35)

p_sih = tf1.paragraphs[0]
p_sih.text = 'SMART INDIA HACKATHON 2026'
p_sih.font.size = Pt(15)
p_sih.font.bold = True
p_sih.font.color.rgb = ACCENT_AMBER

p_main = tf1.add_paragraph()
p_main.text = 'SATYAPAN (सत्यापन — verification)'
p_main.font.size = Pt(28)
p_main.font.bold = True
p_main.font.color.rgb = TEXT_WHITE

p_sub = tf1.add_paragraph()
p_sub.text = 'AI-Based Fake Identity & Document Screening System'
p_sub.font.size = Pt(15)
p_sub.font.color.rgb = ACCENT_CYAN

grid_items = [
    ('▪ Problem Statement ID:', '26188'),
    ('▪ Problem Statement Title:', 'AI-Based Fake Identity & Document Screening System'),
    ('▪ Raised by:', 'Sashastra Seema Bal (SSB), Police II Division'),
    ('▪ Theme:', 'Blockchain & Cybersecurity'),
    ('▪ PS Category:', 'Software'),
    ('▪ Team ID:', '________________'),
    ('▪ Team Name (Registered on portal):', 'Digital Warriors')
]
top_y = Inches(2.3)
for idx, (lbl, val) in enumerate(grid_items):
    b = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.1), top_y + Inches(idx * 0.54), Inches(11.133), Inches(0.46))
    b.fill.solid()
    b.fill.fore_color.rgb = CARD_INNER
    b.line.color.rgb = BORDER_COLOR
    tf_b = b.text_frame
    tf_b.margin_left = Inches(0.2)
    tf_b.margin_top = Inches(0.08)
    p_b = tf_b.paragraphs[0]
    p_b.text = f'{lbl} '
    p_b.font.bold = True
    p_b.font.size = Pt(12)
    p_b.font.color.rgb = ACCENT_BLUE
    run = p_b.add_run()
    run.text = val
    run.font.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = TEXT_WHITE

# SLIDE 2
slide2 = prs.slides.add_slide(blank_layout)
set_bg(slide2)
add_header(slide2, 'SATYAPAN — AI-Based Fake Identity & Document Screening System', slide_num=2)
boxes_s2 = [
    ('Detailed explanation of the proposed solution', [
        '▪ Five modules over one capture: OCR extraction → document validation → tampering detection → face verification → identity resolution',
        '▪ Reads passport, visa, national ID, driving licence and permits — MRZ (ICAO 9303), printed visual zone, and the ePassport RFID chip',
        '▪ Risk engine fuses every signal into CLEAR / REVIEW / REFER with named reasons and a heatmap on the document — in under 3 seconds',
        '▪ Runs on a local GPU box at the counter: offline-capable, no cloud, no traveller data leaving government infrastructure'
    ], ACCENT_BLUE),
    ('How it addresses the problem', [
        '▪ Fake passports and visas → chip signature verification + all five MRZ check digits + print-technology forensics',
        '▪ Altered dates of birth and replaced photographs → MRZ-versus-printed-text mismatch + photo-boundary blur heatmap',
        '▪ Tampered visa stamps → copy-move detection, stamp geometry checks, entry/exit chronology logic',
        '▪ Impersonation and multiple identities → four-way face matching + 1:N vector deduplication across crossings',
        '▪ Expired or blacklisted documents → date and validity rules + locally cached, signed watchlist snapshots',
        '▪ High passenger volume → screening drops from several minutes to a few seconds per traveller'
    ], ACCENT_GREEN),
    ('Innovation and uniqueness of the solution', [
        '▪ Proof, not probability — the passport chip is signed by the issuing state; verifying it makes any altered field fail mathematically, not statistically',
        '▪ Contradiction as evidence — an MRZ-versus-print mismatch is treated as a forgery signal, where other systems smooth it over as an OCR error',
        '▪ Explainable by design — every alert names the check that fired, shows where, and states plainly what could NOT be verified',
        '▪ Catches unknown forgeries — databases only find documents already reported; our AI finds the forgery made last week that is on no list anywhere'
    ], ACCENT_AMBER)
]
for idx, (head, bullets, col) in enumerate(boxes_s2):
    c = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6 + idx * 4.08), Inches(1.35), Inches(3.95), Inches(4.7))
    c.fill.solid()
    c.fill.fore_color.rgb = CARD_BG
    c.line.color.rgb = col
    c.line.width = Pt(1.5)
    tf_c = c.text_frame
    tf_c.word_wrap = True
    tf_c.margin_left = Inches(0.18)
    tf_c.margin_right = Inches(0.18)
    tf_c.margin_top = Inches(0.15)
    p_h = tf_c.paragraphs[0]
    p_h.text = head
    p_h.font.size = Pt(12.5)
    p_h.font.bold = True
    p_h.font.color.rgb = col
    for b in bullets:
        p_b = tf_c.add_paragraph()
        p_b.text = b
        p_b.font.size = Pt(9.5)
        p_b.font.color.rgb = TEXT_WHITE
        p_b.space_before = Pt(5)

bot_y = Inches(6.15)
lbl_bot = slide2.shapes.add_textbox(Inches(0.6), bot_y - Inches(0.25), Inches(12.133), Inches(0.3))
lbl_bot.text_frame.paragraphs[0].text = 'Output of every screening — one decision, named reasons, full evidence:'
lbl_bot.text_frame.paragraphs[0].font.size = Pt(10)
lbl_bot.text_frame.paragraphs[0].font.bold = True
lbl_bot.text_frame.paragraphs[0].font.color.rgb = TEXT_MUTED

dec_cards = [
    ('CLEAR', 'all checks passed · traveller proceeds', ACCENT_GREEN),
    ('REVIEW', 'secondary inspection with named reasons', ACCENT_AMBER),
    ('REFER', 'hard-fail evidence · chip, checksum or watchlist', ACCENT_RED)
]
for idx, (d_title, d_desc, d_col) in enumerate(dec_cards):
    dc = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6 + idx * 4.08), bot_y, Inches(3.95), Inches(0.8))
    dc.fill.solid()
    dc.fill.fore_color.rgb = d_col
    dc.line.fill.background()
    tf_dc = dc.text_frame
    tf_dc.word_wrap = True
    tf_dc.margin_left = Inches(0.1)
    tf_dc.margin_top = Inches(0.08)
    p1 = tf_dc.paragraphs[0]
    p1.alignment = PP_ALIGN.CENTER
    p1.text = d_title
    p1.font.size = Pt(13)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_WHITE
    p2 = tf_dc.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.text = d_desc
    p2.font.size = Pt(9.5)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_WHITE

# SLIDE 3
slide3 = prs.slides.add_slide(blank_layout)
set_bg(slide3)
add_header(slide3, 'TECHNICAL APPROACH', slide_num=3)
c3_l = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.35), Inches(5.9), Inches(4.3))
c3_l.fill.solid()
c3_l.fill.fore_color.rgb = CARD_BG
c3_l.line.color.rgb = ACCENT_BLUE
tf3_l = c3_l.text_frame
tf3_l.word_wrap = True
tf3_l.margin_left = Inches(0.2)
tf3_l.margin_right = Inches(0.2)
tf3_l.margin_top = Inches(0.15)
p3_lh = tf3_l.paragraphs[0]
p3_lh.text = 'Technologies to be used (e.g. programming languages, frameworks, hardware)'
p3_lh.font.size = Pt(11.5)
p3_lh.font.bold = True
p3_lh.font.color.rgb = ACCENT_CYAN

tech_points = [
    '▪ OCR & vision — PaddleOCR PP-OCRv4, TrOCR, LayoutLMv3, YOLOv8-seg, OpenCV',
    '▪ Tampering — TruFor, DocTamper, MVSS-Net + ELA, DCT double-JPEG, Noiseprint++, SIFT copy-move, ExifTool, pikepdf',
    '▪ Face — InsightFace SCRFD + AdaFace, Silent-Face PAD, differential morphing-attack detection',
    '▪ Chip & crypto — BAC/PACE access, Passive Authentication of the SOD, JMRTD, pymrtd, ICAO PKD trust anchors',
    '▪ Risk & fusion — XGBoost with isotonic calibration, SHAP explanations, rule engine',
    '▪ Platform — Python, FastAPI, PostgreSQL, Milvus, MinIO, ONNX Runtime / TensorRT INT8, React + TypeScript',
    '▪ Hardware — VIS/IR/UV document reader, NFC reader (ACR1252U), RGB+IR camera, edge GPU'
]
for tp in tech_points:
    p = tf3_l.add_paragraph()
    p.text = tp
    p.font.size = Pt(9)
    p.font.color.rgb = TEXT_WHITE
    p.space_before = Pt(3.5)

c3_r = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.35), Inches(5.933), Inches(4.3))
c3_r.fill.solid()
c3_r.fill.fore_color.rgb = CARD_BG
c3_r.line.color.rgb = ACCENT_GREEN
tf3_r = c3_r.text_frame
tf3_r.word_wrap = True
tf3_r.margin_left = Inches(0.2)
tf3_r.margin_right = Inches(0.2)
tf3_r.margin_top = Inches(0.15)
p3_rh = tf3_r.paragraphs[0]
p3_rh.text = 'Methodology and process for implementation'
p3_rh.font.size = Pt(11.5)
p3_rh.font.bold = True
p3_rh.font.color.rgb = ACCENT_GREEN

meth_points = [
    '▪ Capture once, then run all five modules in parallel — nothing waits in a queue behind another module',
    '▪ Deterministic checks first: check digits and chip signature can hard-fail a document before any model is consulted',
    '▪ Soft signals feed a calibrated fusion model, so a risk of 0.7 means seven in ten such cases were fraudulent',
    '▪ Officer console shows decision, top three reasons, heatmap and face comparison on one screen',
    '▪ Every override is logged with a reason and feeds retraining; every decision is signed and hash-chained',
    '▪ Models exported to ONNX and quantised to INT8 to hold the latency budget on an affordable GPU'
]
for mp in meth_points:
    p = tf3_r.add_paragraph()
    p.text = mp
    p.font.size = Pt(9)
    p.font.color.rgb = TEXT_WHITE
    p.space_before = Pt(3.5)

flow_card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(5.8), Inches(12.133), Inches(1.2))
flow_card.fill.solid()
flow_card.fill.fore_color.rgb = CARD_INNER
flow_card.line.color.rgb = BORDER_COLOR

f_items = [
    ('CAPTURE\nreader · camera · NFC', Inches(0.8), ACCENT_BLUE),
    ('➔', Inches(2.6), TEXT_MUTED),
    ('QUALITY GATE\ndewarp · deglare · DPI', Inches(3.0), ACCENT_CYAN),
    ('➔', Inches(5.0), TEXT_MUTED),
    ('PARALLEL MODULES\nM1 OCR · M2 Validation\nM3 Tampering · M4 Face · M5 ID', Inches(5.4), ACCENT_AMBER),
    ('➔', Inches(8.5), TEXT_MUTED),
    ('RISK ENGINE\nhard fails + calibrated fusion', Inches(8.9), ACCENT_AMBER),
    ('➔', Inches(10.7), TEXT_MUTED),
    ('OFFICER CONSOLE\nCLEAR / REVIEW / REFER', Inches(11.0), ACCENT_GREEN)
]
for f_txt, f_x, f_c in f_items:
    tb = slide3.shapes.add_textbox(f_x, Inches(5.85), Inches(1.8) if '➔' not in f_txt else Inches(0.4), Inches(1.0))
    p = tb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = f_txt
    p.font.size = Pt(16) if '➔' in f_txt else Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = f_c

# SLIDE 4
slide4 = prs.slides.add_slide(blank_layout)
set_bg(slide4)
add_header(slide4, 'FEASIBILITY AND VIABILITY', slide_num=4)

c4_l = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.35), Inches(5.9), Inches(4.3))
c4_l.fill.solid()
c4_l.fill.fore_color.rgb = CARD_BG
c4_l.line.color.rgb = ACCENT_BLUE
tf4_l = c4_l.text_frame
tf4_l.word_wrap = True
tf4_l.margin_left = Inches(0.2)
tf4_l.margin_top = Inches(0.15)
p4_lh = tf4_l.paragraphs[0]
p4_lh.text = 'Analysis of the feasibility of the idea'
p4_lh.font.size = Pt(12)
p4_lh.font.bold = True
p4_lh.font.color.rgb = ACCENT_CYAN

feas_points = [
    '▪ Every AI module is buildable now with zero external access — public datasets (MIDV-2020, MIDV-Holo, DocTamper, SIDTD, DocXPand) and released pretrained weights',
    '▪ Chip verification is demonstrable today: a ₹4,000 NFC reader plus a publicly published CSCA master list',
    '▪ Public data usable immediately — UN and OFAC sanctions lists, and Aadhaar secure-QR verified against UIDAI\'s public certificate (no licence required)',
    '▪ Commodity hardware: a single edge GPU per lane, no datacentre',
    '▪ Standards are open and stable — ICAO Doc 9303 fully specifies the MRZ, chip and PKI we rely on',
    '▪ Government lookups (Interpol SLTD, IVFRT, CCTNS, ICAO PKD) sit behind documented adapters with clearly labelled simulators'
]
for fp in feas_points:
    p = tf4_l.add_paragraph()
    p.text = fp
    p.font.size = Pt(9)
    p.font.color.rgb = TEXT_WHITE
    p.space_before = Pt(3.5)

c4_r = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.35), Inches(5.933), Inches(4.3))
c4_r.fill.solid()
c4_r.fill.fore_color.rgb = CARD_BG
c4_r.line.color.rgb = ACCENT_AMBER
tf4_r = c4_r.text_frame
tf4_r.word_wrap = True
tf4_r.margin_left = Inches(0.2)
tf4_r.margin_top = Inches(0.15)
p4_rh = tf4_r.paragraphs[0]
p4_rh.text = 'Potential challenges and risks → Strategies for overcoming them'
p4_rh.font.size = Pt(11.5)
p4_rh.font.bold = True
p4_rh.font.color.rgb = ACCENT_AMBER

chal_points = [
    '▪ False alarms erode officer trust → Tune on false-positive rate at fixed detection rate; shadow-mode pilot before live decisions.',
    '▪ No public training data for Nepali/Bhutanese docs → Programmatic forgery generation over synthetic templates only under SSB authorization.',
    '▪ Forgers adapt to deployed model → Signal diversity (font, print-tech, UV checks); adversarial retraining and rapid rollback.',
    '▪ Demographic bias in face matching → Stratified error reporting by nationality, sex, age; calibrated per NIST FRVT practice.',
    '▪ Data protection under DPDP & Aadhaar Act → On-premise only, Aadhaar masked, keys in HSM, hash-chained audit log.',
    '▪ Visible-light photos limit detection → Software works on ordinary images; improves with VIS/IR/UV reader + chip verification.'
]
for cp in chal_points:
    p = tf4_r.add_paragraph()
    p.text = cp
    p.font.size = Pt(8.8)
    p.font.color.rgb = TEXT_WHITE
    p.space_before = Pt(3.5)

phase_card = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(5.8), Inches(12.133), Inches(1.2))
phase_card.fill.solid()
phase_card.fill.fore_color.rgb = CARD_INNER
phase_card.line.color.rgb = BORDER_COLOR

phases = [
    ('PHASE 1 · Internal Round', 'MRZ pipeline · check-digit validation · rule engine · classical tampering ensemble · face matching · officer console · mocked adapters', ACCENT_BLUE, Inches(0.8)),
    ('PHASE 2 · Grand Finale', 'ePassport chip Passive Authentication · synthetic forgery training · 1:N deduplication · calibrated fusion with SHAP · ONNX optimisation', ACCENT_CYAN, Inches(4.8)),
    ('PHASE 3 · With Department', 'VIS/IR/UV reader · ICAO PKD and Interpol SLTD under MHA authorisation · CERT-In audit · shadow-mode pilot at live checkpoint', ACCENT_GREEN, Inches(8.8))
]
for p_t, p_d, p_c, p_x in phases:
    tb = slide4.shapes.add_textbox(p_x, Inches(5.85), Inches(3.8), Inches(1.0))
    p1 = tb.text_frame.paragraphs[0]
    p1.text = p_t
    p1.font.size = Pt(11)
    p1.font.bold = True
    p1.font.color.rgb = p_c
    p2 = tb.text_frame.add_paragraph()
    p2.text = p_d
    p2.font.size = Pt(8.5)
    p2.font.color.rgb = TEXT_WHITE

# SLIDE 5
slide5 = prs.slides.add_slide(blank_layout)
set_bg(slide5)
add_header(slide5, 'IMPACT AND BENEFITS', slide_num=5)

kpi_data = [
    ('< 3 sec', 'per document, versus several minutes manually', ACCENT_CYAN),
    ('4', 'independent face comparisons per traveller', ACCENT_GREEN),
    ('100%', 'of decisions logged with a signed, tamper-evident trail', ACCENT_AMBER)
]
for idx, (k_val, k_lbl, k_col) in enumerate(kpi_data):
    kc = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6 + idx * 4.08), Inches(1.35), Inches(3.95), Inches(0.95))
    kc.fill.solid()
    kc.fill.fore_color.rgb = CARD_INNER
    kc.line.color.rgb = k_col
    tf_kc = kc.text_frame
    tf_kc.word_wrap = True
    tf_kc.margin_left = Inches(0.15)
    tf_kc.margin_top = Inches(0.08)
    p1 = tf_kc.paragraphs[0]
    p1.text = k_val
    p1.font.size = Pt(22)
    p1.font.bold = True
    p1.font.color.rgb = k_col
    p2 = tf_kc.add_paragraph()
    p2.text = k_lbl
    p2.font.size = Pt(9.5)
    p2.font.color.rgb = TEXT_WHITE

c5_l = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(2.45), Inches(5.9), Inches(2.9))
c5_l.fill.solid()
c5_l.fill.fore_color.rgb = CARD_BG
c5_l.line.color.rgb = ACCENT_BLUE
tf5_l = c5_l.text_frame
tf5_l.word_wrap = True
tf5_l.margin_left = Inches(0.2)
tf5_l.margin_top = Inches(0.12)
p5_lh = tf5_l.paragraphs[0]
p5_lh.text = 'Potential impact on the target audience'
p5_lh.font.size = Pt(12)
p5_lh.font.bold = True
p5_lh.font.color.rgb = ACCENT_CYAN

imp_points = [
    '▪ Border officers (SSB, Bureau of Immigration) — screening collapses from minutes to seconds; decisions rest on evidence rather than shift fatigue',
    '▪ Travellers — genuine documents clear quickly, queues shorten; flagged travellers get named reasons and documented route to redress',
    '▪ Investigators and prosecutors — signed, hash-chained record per crossing with Section 63 electronic-evidence certificate generated automatically',
    '▪ Ministry of Home Affairs — same forgery scores identical at every checkpoint, creating national verification standard'
]
for ip in imp_points:
    p = tf5_l.add_paragraph()
    p.text = ip
    p.font.size = Pt(9)
    p.font.color.rgb = TEXT_WHITE
    p.space_before = Pt(3)

c5_r = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(2.45), Inches(5.933), Inches(2.9))
c5_r.fill.solid()
c5_r.fill.fore_color.rgb = CARD_BG
c5_r.line.color.rgb = ACCENT_GREEN
tf5_r = c5_r.text_frame
tf5_r.word_wrap = True
tf5_r.margin_left = Inches(0.2)
tf5_r.margin_top = Inches(0.12)
p5_rh = tf5_r.paragraphs[0]
p5_rh.text = 'Benefits of the solution (social, economic, environmental)'
p5_rh.font.size = Pt(12)
p5_rh.font.bold = True
p5_rh.font.color.rgb = ACCENT_GREEN

ben_points = [
    '▪ Security — detects tampering/impersonation without watchlists, including morphed photos',
    '▪ Social — consistent non-arbitrary counter treatment; error rates measured across demographics',
    '▪ Economic — higher throughput without adding posts; commodity hardware avoids costly SDK licensing',
    '▪ Operational & Environmental — offline-first at low-connectivity posts; digital evidence eliminates paper case files'
]
for bp in ben_points:
    p = tf5_r.add_paragraph()
    p.text = bp
    p.font.size = Pt(9)
    p.font.color.rgb = TEXT_WHITE
    p.space_before = Pt(3)

c5_cmp_l = slide5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(5.5), Inches(5.9), Inches(1.55))
c5_cmp_l.fill.solid()
c5_cmp_l.fill.fore_color.rgb = RGBColor(30, 20, 20)
c5_cmp_l.line.color.rgb = ACCENT_RED
tf5_cl = c5_cmp_l.text_frame
tf5_cl.word_wrap = True
tf5_cl.margin_left = Inches(0.15)
tf5_cl.margin_top = Inches(0.08)
p_cl = tf5_cl.paragraphs[0]
p_cl.text = 'TODAY — MANUAL INSPECTION'
p_cl.font.size = Pt(11)
p_cl.font.bold = True
p_cl.font.color.rgb = ACCENT_RED
t_pts = ['• Several minutes per document at the counter', '• Depends on officer shift fatigue and queue pressure', '• Sophisticated forgeries routinely missed by eye', '• No audit record of why traveller was cleared or stopped']
for pt in t_pts:
    p = tf5_cl.add_paragraph()
    p.text = pt
    p.font.size = Pt(8.5)
    p.font.color.rgb = TEXT_MUTED

c5_cmp_r = slide5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.8), Inches(5.5), Inches(5.933), Inches(1.55))
c5_cmp_r.fill.solid()
c5_cmp_r.fill.fore_color.rgb = RGBColor(15, 30, 25)
c5_cmp_r.line.color.rgb = ACCENT_GREEN
tf5_cr = c5_cmp_r.text_frame
tf5_cr.word_wrap = True
tf5_cr.margin_left = Inches(0.15)
tf5_cr.margin_top = Inches(0.08)
p_cr = tf5_cr.paragraphs[0]
p_cr.text = 'WITH SATYAPAN'
p_cr.font.size = Pt(11)
p_cr.font.bold = True
p_cr.font.color.rgb = ACCENT_GREEN
s_pts = ['• Under three seconds, with the officer still deciding', '• The same forgery scores identically at every checkpoint', '• Chip-verified proof plus multi-signal explainable AI ensemble', '• Signed, hash-chained electronic evidence trail for every decision']
for pt in s_pts:
    p = tf5_cr.add_paragraph()
    p.text = pt
    p.font.size = Pt(8.5)
    p.font.color.rgb = TEXT_WHITE

# SLIDE 6
slide6 = prs.slides.add_slide(blank_layout)
set_bg(slide6)
add_header(slide6, 'RESEARCH AND REFERENCES', slide_num=6)

ref_boxes = [
    ('Standards & Specifications', [
        '▪ ICAO Doc 9303, Parts 3, 4, 10, 11 and 12 — machine readable travel documents, logical data structure and PKI (icao.int)',
        '▪ ICAO Public Key Directory — CSCA and document signer trust anchors (icao.int/Security/FAL/PKD)',
        '▪ ISO/IEC 30107-3 — presentation attack detection reporting; ISO/IEC 19794-5 and 39794-5 — face image data'
    ], ACCENT_BLUE),
    ('Models & Methods', [
        '▪ TruFor: Leveraging All-Round Clues for Trustworthy Image Forgery Detection and Localization — CVPR 2023',
        '▪ CAT-Net: Compression Artifact Tracing Network for splicing localisation — IJCV 2022',
        '▪ MVSS-Net++: multi-view multi-scale supervision for image manipulation detection — TPAMI 2022',
        '▪ AdaFace: Quality Adaptive Margin for Face Recognition — CVPR 2022; ArcFace — CVPR 2019',
        '▪ NIST FRVT Part 4 (MORPH) — morphing attack detection benchmark (nist.gov)'
    ], ACCENT_GREEN),
    ('Datasets Referenced', [
        '▪ MIDV-500 / MIDV-2019 / MIDV-2020 — identity document images and video, Institute for Systems Analysis, RAS',
        '▪ MIDV-Holo — holographic security feature forgery detection',
        '▪ DocXPand-25k and SIDTD — synthetic and tampered identity documents',
        '▪ DocTamper (2023) — document-specific tampering localisation',
        '▪ FRLL-Morphs, SMDD — face morphing benchmarks; CASIA-SURF, OULU-NPU — presentation attack',
        '▪ CASIA v2, IMD2020, DEFACTO, CoMoFoD — general image manipulation'
    ], ACCENT_CYAN),
    ('Indian Systems, Legal Framework & Hardware', [
        '▪ API Setu — NIC platform for government issuer APIs including Sarathi and Vahan (apisetu.gov.in)',
        '▪ DigiLocker Partner APIs — issued document retrieval with citizen consent (partners.digitallocker.gov.in)',
        '▪ UIDAI — Aadhaar secure QR code and paperless offline e-KYC specifications (uidai.gov.in)',
        '▪ Legal: Digital Personal Data Protection Act 2023 Section 17(2)(a), Aadhaar Act 2016 Section 29, BSA 2023 Section 63',
        '▪ Hardware: ACS ACR1252U NFC reader · Regula/Access-IS VIS/IR/UV readers · Intel RealSense · NVIDIA Jetson AGX Orin'
    ], ACCENT_AMBER)
]

for idx, (r_head, r_items, r_col) in enumerate(ref_boxes):
    col_idx = idx % 2
    row_idx = idx // 2
    x = Inches(0.6 + col_idx * 6.2)
    y = Inches(1.35 + row_idx * 2.85)
    
    rc = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.933), Inches(2.7))
    rc.fill.solid()
    rc.fill.fore_color.rgb = CARD_BG
    rc.line.color.rgb = r_col
    rc.line.width = Pt(1.5)
    tf_rc = rc.text_frame
    tf_rc.word_wrap = True
    tf_rc.margin_left = Inches(0.18)
    tf_rc.margin_right = Inches(0.18)
    tf_rc.margin_top = Inches(0.1)
    
    pr = tf_rc.paragraphs[0]
    pr.text = r_head
    pr.font.size = Pt(11.5)
    pr.font.bold = True
    pr.font.color.rgb = r_col
    
    for item in r_items:
        p = tf_rc.add_paragraph()
        p.text = item
        p.font.size = Pt(8.5)
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(3)

out_file = r'E:\SIH\DocShield_Bharat_AI_SIH2026_PPT.pptx'
prs.save(out_file)
print('SUCCESSFULLY_CREATED_EXACT_PPTX')
